from io import BytesIO
from pptx import Presentation
import tiktoken
from core.config import settings

def pptx_text_extraction(blob_data):
    document_text = ""
    # Create an in-memory binary stream from the blob data
    pptx_stream = BytesIO(blob_data)
    # Load the presentation from the binary stream
    presentation = Presentation(pptx_stream)
    # Iterate through slides and shapes to extract text
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                document_text = document_text + shape.text + '\n'
    return document_text

def get_num_tokens(text:str):
    # Get encoding
    encoding = tiktoken.get_encoding(settings.TIKTOKEN_ENCODING_NAME)
    # Pass the text and get encoded tokens.
    encoded_tokens = encoding.encode(text)
    # Count number of tokens
    num_tokens = len(encoded_tokens)
    
    return num_tokens